"""Build the editable FreshFind project report from the inspected application."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "FreshFind_Project_Report.docx"
IMAGES = ROOT / "docs" / "report-assets"
GREEN = RGBColor(31, 75, 55)
OLIVE = RGBColor(96, 117, 74)
MUTED = RGBColor(89, 100, 85)
WHITE = RGBColor(255, 255, 255)
PALE = "E8EDE0"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, value, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(value)
    run.bold = bold
    run.font.size = Pt(8.4)
    run.font.color.rgb = color or GREEN
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Light Shading Accent 1"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, True, WHITE)
        shade(table.rows[0].cells[i], "1F4B37")
        if widths:
            table.rows[0].cells[i].width = Inches(widths[i])
    for index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            set_cell_text(cells[i], str(value))
            if widths:
                cells[i].width = Inches(widths[i])
            if index % 2 == 0:
                shade(cells[i], "F2F5ED")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def para(doc, text="", style=None):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def figure(doc, filename, caption):
    image = IMAGES / filename
    if image.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        width = 5.8 if filename in {"market-branch.png", "market-directory.png"} else 6.7
        p.add_run().add_picture(str(image), width=Inches(width))
        cap = para(doc, caption, "Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def make_diagrams():
    IMAGES.mkdir(exist_ok=True)
    font_path = "C:/Windows/Fonts/arial.ttf"
    font = ImageFont.truetype(font_path, 28)
    small = ImageFont.truetype(font_path, 23)
    green = "#1f4b37"
    pale = "#eef2e7"

    def box(draw, coordinates, label, tone=pale):
        draw.rounded_rectangle(coordinates, radius=22, fill=tone, outline=green, width=3)
        cx = (coordinates[0] + coordinates[2]) / 2
        cy = (coordinates[1] + coordinates[3]) / 2
        draw.multiline_text((cx, cy), label, font=font, fill=green, anchor="mm", align="center", spacing=5)

    def arrow(draw, a, b):
        draw.line((a, b), fill=green, width=5)
        x, y = b
        draw.polygon([(x, y), (x - 19, y - 10), (x - 19, y + 10)], fill=green)

    flow = Image.new("RGB", (1600, 260), "white")
    d = ImageDraw.Draw(flow)
    steps = [(20, 75, 270, 185, "Open\ndirectory"), (335, 75, 585, 185, "Search or\nfilter"), (650, 75, 900, 185, "Select\nmarket"), (965, 75, 1215, 185, "View details\nand schedule"), (1280, 75, 1530, 185, "Save or get\ndirections")]
    for x1, y1, x2, y2, label in steps:
        box(d, (x1, y1, x2, y2), label)
    for x in [280, 595, 910, 1225]:
        arrow(d, (x, 130), (x + 43, 130))
    flow.save(IMAGES / "market-flow.png")

    dfd = Image.new("RGB", (1600, 390), "white")
    d = ImageDraw.Draw(dfd)
    box(d, (25, 70, 350, 185), "Visitor\n(browser)")
    box(d, (595, 70, 1005, 185), "FreshFind\napplication", "#dfead6")
    box(d, (1245, 70, 1575, 185), "Bundled JSON\nmarket and produce")
    box(d, (380, 265, 705, 370), "Browser\nlocal storage")
    box(d, (900, 265, 1225, 370), "Google Maps\nexternal link")
    arrow(d, (360, 115), (580, 115))
    arrow(d, (1015, 115), (1230, 115))
    d.line((595, 170, 455, 265), fill=green, width=5)
    d.polygon([(455, 265), (475, 245), (484, 263)], fill=green)
    d.line((1005, 170, 1080, 265), fill=green, width=5)
    d.polygon([(1080, 265), (1055, 254), (1075, 240)], fill=green)
    d.text((450, 25), "queries and results", font=small, fill=green)
    d.text((1040, 25), "read static data", font=small, fill=green)
    dfd.save(IMAGES / "data-flow.png")


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(.72)
    section.bottom_margin = Inches(.72)
    section.left_margin = Inches(.78)
    section.right_margin = Inches(.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = GREEN
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.18
    for level, size in [(1, 17), (2, 12.5), (3, 10.5)]:
        style = styles[f"Heading {level}"]
        style.font.name = "Georgia"
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = GREEN
        style.paragraph_format.space_before = Pt(13 if level == 1 else 9)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    styles["Title"].font.name = "Georgia"
    styles["Title"].font.size = Pt(30)
    styles["Title"].font.bold = False
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)
    styles["Subtitle"].font.name = "Aptos"
    styles["Subtitle"].font.size = Pt(12)
    styles["Subtitle"].font.color.rgb = MUTED
    styles["Caption"].font.name = "Aptos"
    styles["Caption"].font.size = Pt(8)
    styles["Caption"].font.italic = True
    styles["Caption"].font.color.rgb = MUTED
    styles["Caption"].paragraph_format.space_after = Pt(9)

    make_diagrams()

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("FreshFind  |  Project report  |  September 2026").font.size = Pt(8)

    para(doc, "FRESHFIND  /  FRESH ALL ALONG").runs[0].font.color.rgb = OLIVE
    title = doc.add_paragraph(style="Title")
    title.add_run("FreshFind Website Project Report")
    para(doc, "Requirements, design, test data and installation", "Subtitle")
    para(doc, "Version 1.0  |  25 September 2026")
    para(doc, "FreshFind is a browser based guide to selected Lagos farmers markets and seasonal produce. This report records the features present in the current React application, explains how its static data and rule based assistant work, and compares the implementation with the supplied FreshFind software requirements specification. It also gives test scenarios and installation instructions for review or handover.")
    add_table(doc, ["Document basis", "Role in this report"], [
        ("FreshFind Web Innovation Unleashed SRS, version 1.0", "Requirements baseline and required report sections"),
        ("FreshFind repository inspected 25 September 2026", "Evidence for current features, structure and test data"),
        ("CarBreezy project report example", "General report structure and presentation reference only"),
    ], [2.65, 4.05])
    heading(doc, "Contents")
    for item in ["1  Problem Definition and Project Scope", "2  Implemented Website Features", "3  Design Specifications and Architecture", "4  User Flows and Data Flow", "5  Data Model and Test Data", "6  Verification and Quality", "7  Installation and Operation", "8  Requirements Remaining and Handover"]:
        para(doc, item)

    doc.add_page_break()
    heading(doc, "1  Problem Definition and Project Scope")
    para(doc, "Residents often learn about local market days, hours and available produce through scattered posts or word of mouth. FreshFind brings a selected set of Lagos market details into one website so a visitor can discover a market, check its regular schedule and explore produce before planning a trip. The current dataset is illustrative, so the site advises visitors to confirm schedules directly with a market.")
    para(doc, "The application is a single page React website. It runs in the browser, reads bundled JSON data, and has no application server or live market database. Market and produce listings do not update from the website itself. Saved items are kept in the visitor's browser storage.")
    heading(doc, "Project objectives", 2)
    for item in ["Help visitors search markets by name, area or listed produce.", "Show regular opening days and hours, market locations and typical items.", "Give visitors a searchable guide to produce and recorded seasons.", "Provide a rule based assistant for common market and produce questions.", "Let visitors save markets and produce for later on the same browser."]:
        para(doc, item, "List Bullet")
    heading(doc, "2  Implemented Website Features")
    add_table(doc, ["Area", "Current behaviour"], [
        ("Home", "Tree themed landing page with canopy, three scroll driven branches for markets, produce and saved items, roots and soil footer. Search opens the market directory."),
        ("Market directory", "Lists five markets. Free text searches name, area, address and listed produce. Area and open today controls filter the list."),
        ("Market detail", "Shows description, regular days and hours, produce, current open status and a Google Maps directions link."),
        ("Produce guide", "Lists eight produce entries. Search, category and current season controls narrow results."),
        ("Saved items", "Visitors save or remove markets and produce. The collection persists in browser local storage."),
        ("Assistant", "Floating widget accepts typed questions and quick prompts; responses come from local JSON and fixed rules."),
    ], [1.3, 5.4])
    heading(doc, "Current website screens", 2)
    figure(doc, "market-branch.png", "Figure 1  Home page market branch attached to the tree trunk")
    figure(doc, "market-directory.png", "Figure 2  Searchable Lagos market directory")

    heading(doc, "3  Design Specifications and Architecture")
    para(doc, "React 19 renders the interface and React Router maps browser paths to pages. Vite 8 builds the project. CSS handles the responsive layout and tree animation; Lucide supplies icons. The home page rotates its decorative trunk and section panels from scroll position. Animation can be paused on desktop and is disabled for reduced motion or narrow screens. The earlier home concepts remain available in the repository for reference.")
    heading(doc, "Site map", 2)
    add_table(doc, ["Path", "Purpose"], [
        ("/", "Tree themed landing page"),
        ("/markets", "Market search and filters"),
        ("/markets/:id", "Individual market details"),
        ("/produce-guide", "Produce catalogue"),
        ("/saved", "Visitor's saved collection"),
        ("/h and /h2", "Earlier home page prototypes retained for comparison"),
    ], [1.35, 5.35])
    heading(doc, "Application components", 2)
    add_table(doc, ["Layer", "Responsibility"], [
        ("Page and navigation", "App routes, Navbar and Footer display the active page and shared shell."),
        ("Feature components", "MarketCard, ProduceCard and ChatBot render reusable visitor actions."),
        ("Local data", "JSON files supply markets, produce, seasonal entries and assistant responses."),
        ("Browser state", "SavedContext reads and writes localStorage key freshfind-saved."),
        ("External navigation", "Market details open Google Maps search links for coordinates; this is not an embedded map."),
    ], [1.45, 5.25])

    heading(doc, "4  User Flows and Data Flow")
    heading(doc, "Market discovery flow", 2)
    add_table(doc, ["Step", "Visitor action", "Website response"], [
        ("1", "Open the home page or market directory", "Display the tree introduction or all market cards."),
        ("2", "Enter a market, area or produce term", "Filter local market data and update the result count."),
        ("3", "Choose an area or open today", "Narrow the cards using area and Lagos market day rules."),
        ("4", "Open a market card", "Show schedule, listed produce, location and directions link."),
        ("5", "Save the market", "Add its ID to localStorage and show it under Saved."),
    ], [.55, 2.05, 4.1])
    figure(doc, "market-flow.png", "Figure 3  Market discovery flowchart")
    heading(doc, "Level 0 data flow", 2)
    figure(doc, "data-flow.png", "Figure 4  Level 0 data flow diagram")
    para(doc, "The application matches the request against static data, renders cards or assistant replies, and returns results to the visitor. A save action writes only to that visitor's browser localStorage. A directions action opens Google Maps in a new browser tab. No information is sent to a FreshFind server.")
    heading(doc, "Assistant decision flow", 2)
    para(doc, "Question or quick prompt  >  text normalization  >  produce, market, area, day or season rule  >  matching records and reply card  >  fallback reply if no rule matches")
    para(doc, "The assistant is rule based and does not connect to an external AI service. Its open today and season rules use the browser's local date; market card open status uses the Africa/Lagos time zone. This difference matters when a visitor browses from another time zone.")

    heading(doc, "5  Data Model and Test Data")
    para(doc, "The project stores five example Lagos markets and eight produce entries in prepopulated JSON files. The market dataset contains ID, name, area, address, description, regular days, opening and closing times, typical produce and coordinates. A produce record contains ID, name, category, description and named season months. Assistant responses and seasonal prompts have separate local JSON files.")
    add_table(doc, ["Dataset", "Records used", "Sample values"], [
        ("Markets", "5", "Ikeja, Lekki, Yaba, Surulere, Victoria Island"),
        ("Produce", "8", "Tomatoes, Carrots, Spinach, Bananas, Yam, Pepper, Cucumber, Watermelon"),
        ("Saved IDs", "Per browser", "Two arrays named markets and produce in localStorage"),
    ], [1.25, 1.0, 4.45])
    heading(doc, "Representative test scenarios", 2)
    add_table(doc, ["Scenario", "Input", "Expected result"], [
        ("Market search", "Ikeja", "Ikeja Fresh Market appears."),
        ("Produce search in markets", "Tomatoes", "Markets listing Tomatoes remain."),
        ("Area filter", "Lekki", "Only the Lekki market card appears."),
        ("Produce category", "Fruits", "Bananas and Watermelon remain."),
        ("Bookmark", "Save Ikeja market", "Card is marked saved and appears under Saved markets after reload."),
        ("Assistant", "Where can I find tomatoes?", "Matching markets appear as linked reply cards."),
        ("Empty results", "Unlisted search term", "A no results message is shown."),
    ], [1.3, 1.75, 3.65])

    heading(doc, "6  Verification and Quality")
    para(doc, "The project completed a production Vite build and ESLint check during this documentation pass. Browser inspection confirmed the home page, market branch and market directory render with expected navigation and content at a desktop viewport. These checks are limited to the inspected environment; a full cross browser, accessibility and Lighthouse audit is still required before formal acceptance.")
    add_table(doc, ["Check", "Observed result", "Further validation"], [
        ("npm run build", "Passed", "Repeat after data or UI changes."),
        ("npm run lint", "Passed", "Repeat after code changes."),
        ("Desktop browser inspection", "Home and directory rendered", "Test all paths and controls on mobile, keyboard and other browsers."),
        ("Lighthouse", "Not run", "Record performance, accessibility and SEO scores."),
    ], [1.7, 1.9, 3.1])
    para(doc, "Accessibility provisions visible in the implementation include labelled search inputs, a skip link, keyboard focus styles, descriptive buttons, reduced motion handling and alternate text for content images. The chat widget returns focus to its launcher when closed. These provisions should be checked with keyboard and screen reader testing rather than treated as certification.")

    heading(doc, "7  Installation and Operation")
    para(doc, "Prerequisites: a recent Node.js release compatible with Vite 8, npm and a current browser. The repository contains a package lock file for repeatable dependency installation. No database, API key or backend service is required for the current website.")
    add_table(doc, ["Step", "Command or action", "Result"], [
        ("1", "Open the FreshFind repository in a terminal", "The project folder is the working directory."),
        ("2", "npm ci", "Install dependencies from package lock."),
        ("3", "npm run dev", "Start the Vite local development server; open the URL it prints."),
        ("4", "npm run build", "Create a production build in dist."),
        ("5", "npm run preview", "Serve that build locally for review."),
        ("6", "npm run lint", "Run the repository's ESLint check."),
    ], [.55, 1.65, 4.5])
    para(doc, "For static hosting, publish the contents of dist and configure the host to serve index.html for client side routes such as /markets/1. Market data lives in src/JSON; edit those files, check references between market produce names and guide entries, then rebuild. Browser bookmarks are local to each visitor and will not follow them to another device.")

    heading(doc, "8  Requirements Remaining and Handover")
    para(doc, "The SRS is broader than the current implementation. The following items should be completed or explicitly waived before calling the full specification delivered.")
    add_table(doc, ["SRS item", "Current state", "Work remaining"], [
        ("Directory day and sort controls", "Partial", "Add a chosen weekday filter and alphabetic, proximity or next open sorting."),
        ("Embedded map", "Partial", "Market pages link to Google Maps but do not embed a map."),
        ("Bookmark notes, export and sharing", "Missing", "Add notes and export or share actions if required."),
        ("About and Contact pages", "Missing from active routes", "Build and link both pages, or obtain a scope change."),
        ("Visitor counter, clock and geolocation", "Missing from active UI", "Implement only if retained as acceptance criteria."),
        ("Login and signup mock controls", "Missing", "Add the requested design only controls if required."),
        ("Lighthouse and full browser testing", "Not recorded", "Run and document the acceptance checks."),
        ("Demonstration video and source ZIP", "Not part of website code", "Prepare the SRS submission artifacts separately."),
    ], [2.0, 1.35, 3.35])
    heading(doc, "Source references", 2)
    para(doc, "FreshFind Web Innovation Unleashed Software Requirements Specification, version 1.0, supplied PDF. FreshFind source repository and bundled JSON data inspected 25 September 2026. CarBreezy project report supplied as an example of report organization; its project claims and code were not used as FreshFind evidence.")

    OUTPUT.parent.mkdir(exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
